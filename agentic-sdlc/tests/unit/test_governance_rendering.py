"""Enhancement 3: the shared DecisionPackage renderer (governance/
rendering.py) used by `agentic approvals show`, the watch GUI panel, and
`agentic approvals export --format docx`. One content model
(render_sections), two output formats (render_text, render_docx) — these
tests build a package directly rather than running a full workflow, so
they exercise the renderer in isolation.
"""

from datetime import UTC, datetime
from pathlib import Path
from typing import Literal

from agentic.core.models import Artifact, Decision, PolicyVerdict, PolicyVerdictType
from agentic.governance.approvals import DecisionPackage
from agentic.governance.rendering import Table, render_docx, render_sections, render_text

ArtifactKind = Literal["spec", "design", "code", "test", "doc", "report", "analysis"]


def _artifact(name: str, payload: dict | str, kind: ArtifactKind = "design") -> Artifact:
    return Artifact(
        artifact_id=f"art-{name}", name=name, kind=kind, content_hash="hash",
        payload=payload, produced_by_node="design.arch", produced_by_agent="architect",
        run_id="run-1", created_at=datetime.now(UTC),
    )


def _package(artifacts: list[Artifact], decisions: list[Decision] | None = None,
             policy_verdicts: list[PolicyVerdict] | None = None) -> DecisionPackage:
    return DecisionPackage(
        node_id="design.review", stage="DESIGN_REVIEW", input_hash="deadbeef",
        artifacts=artifacts, decisions=decisions or [], policy_verdicts=policy_verdicts or [],
        blast_radius={"files": [], "endpoints": []},
        consequence_of_rejection="the node is routed to rollback",
    )


def test_named_sections_render_in_declared_order_with_tables_for_structured_data() -> None:
    package = _package([
        _artifact("normalized_spec", {"text": "Build a widget."}, kind="spec"),
        _artifact("task_graph", {"tasks": [
            {"task_id": "T1", "description": "Do the thing", "depends_on": [], "effort": "M", "risk": "low"},
        ]}, kind="spec"),
        _artifact("architecture_design", {
            "components": ["api", "service"], "boundaries": "layered", "cross_cutting_concerns": ["logging"],
        }),
        _artifact("adr_records", {"decisions": [
            {"statement": "Use layers", "rationale": "separation of concerns", "alternatives_rejected": ["monolith"]},
        ]}),
        _artifact("data_model_design", {"entities": [
            {"name": "Widget", "fields": [{"name": "id", "type": "str", "constraints": "uuid"}]},
        ]}),
        _artifact("openapi_schema", {"endpoints": [
            {"method": "POST", "path": "/widgets", "summary": "Create a widget"},
        ]}),
    ])

    sections = render_sections(package)
    titles = [s.title for s in sections]

    assert titles[:6] == [
        "Normalized Requirement", "Task Decomposition", "Architecture Design",
        "Architecture Decision Records", "Data Model", "API Contract",
    ]
    assert titles[-1] == "Checkpoint"

    task_section = next(s for s in sections if s.title == "Task Decomposition")
    table = next(b for b in task_section.blocks if isinstance(b, Table))
    assert table.headers == ["Task", "Description", "Depends on", "Effort", "Risk"]
    assert table.rows == [["T1", "Do the thing", "", "M", "low"]]


def test_decisions_and_policy_verdicts_get_their_own_sections() -> None:
    package = _package(
        [],
        decisions=[
            Decision(
                decision_id="d1", node_id="design.arch", agent="architect",
                statement="Use layers", rationale="separation of concerns", created_at=datetime.now(UTC),
            )
        ],
        policy_verdicts=[PolicyVerdict(rule_id="R1", category="security", verdict=PolicyVerdictType.WARN, message="check this")],
    )
    sections = render_sections(package)
    titles = [s.title for s in sections]
    assert "Decisions" in titles
    assert "Policy Evaluations" in titles

    decisions_table = next(b for b in next(s for s in sections if s.title == "Decisions").blocks if isinstance(b, Table))
    assert decisions_table.rows == [["architect", "Use layers", "separation of concerns"]]


def test_generic_fallback_avoids_raw_dict_repr_for_unnamed_artifacts() -> None:
    package = _package([
        _artifact("ambiguity_assessment", {
            "scored": [{"id": "AMB-1", "question": "Real db?", "proposed_default": "no"}],
            "assumptions": ["in-memory is fine", "no rate limiting"],
        }, kind="analysis"),
        _artifact("clarification_questions", {"questions": []}, kind="analysis"),
    ])
    rendered = render_text(package)

    assert "{'id':" not in rendered
    assert "[{'" not in rendered
    assert "in-memory is fine, no rate limiting" in rendered
    assert "(none)" in rendered  # empty "questions" list


def test_render_text_has_no_table_rows_for_empty_task_list() -> None:
    package = _package([_artifact("task_graph", {"tasks": []}, kind="spec")])
    rendered = render_text(package)
    assert "(no tasks)" in rendered


def test_render_docx_produces_headings_and_real_tables(tmp_path: Path) -> None:
    from docx import Document

    package = _package([
        _artifact("normalized_spec", {"text": "Build a widget."}, kind="spec"),
        _artifact("openapi_schema", {"endpoints": [
            {"method": "GET", "path": "/widgets/{id}", "summary": "Get a widget"},
        ]}),
    ])
    out_path = tmp_path / "package.docx"
    render_docx(package, out_path)

    assert out_path.exists()
    doc = Document(str(out_path))
    heading_texts = [p.text for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")]
    assert "Normalized Requirement" in heading_texts
    assert "API Contract" in heading_texts
    assert len(doc.tables) >= 1
    assert doc.tables[0].rows[0].cells[0].text == "Method"


def test_release_manager_output_gets_its_own_readable_sections() -> None:
    """The actual bug report: `approvals show ... release.readiness` must
    surface the release_manager's own output — go/no-go, risks,
    limitations, plus upstream test/coverage/security results — not just
    the same design-review artifacts every other gate shows."""
    package = _package([
        _artifact("test_results", {"exit_code": 0, "coverage": 0.96}, kind="report"),
        _artifact("security_findings", {"findings": [
            {"severity": "LOW", "description": "in-memory store", "location": "app/repository.py"},
        ]}, kind="report"),
        _artifact("release_readiness_report", {
            "go_no_go": "go", "risks": ["no rate limiting"], "limitations": ["in-memory persistence"],
        }, kind="report"),
        _artifact("engineering_summary", {"summary": "All exit gates passed."}, kind="report"),
    ])
    rendered = render_text(package)

    assert "Unit Test Results" in rendered
    assert "Coverage: 96%" in rendered
    assert "Security Findings" in rendered
    assert "LOW" in rendered
    assert "Release Readiness" in rendered
    assert "Go/No-Go: GO" in rendered
    assert "no rate limiting" in rendered
    assert "Engineering Summary" in rendered
    assert "All exit gates passed." in rendered


def test_generated_source_files_are_summarized_not_dumped_in_full() -> None:
    """code/test/doc artifacts hold the entire generated file as a raw
    string payload — a decision package must not become a multi-hundred-
    line source dump; the workspace/git diff is where full content
    belongs, not a human-readable review."""
    long_source = "\n".join(f"line {i}" for i in range(200))
    package = _package([
        _artifact("code:app/main.py", long_source, kind="code"),
        _artifact("test:tests/test_api.py", long_source, kind="test"),
        _artifact("doc:README.md", "short readme", kind="doc"),
    ])
    rendered = render_text(package)

    assert "line 199" not in rendered
    assert "200 lines" in rendered
    assert "see workspace for full content" in rendered


def test_render_text_and_render_docx_agree_on_section_titles(tmp_path: Path) -> None:
    """The single-source-of-truth contract: both formats must render
    exactly the sections render_sections() produces, in the same order."""
    from docx import Document

    package = _package([
        _artifact("normalized_spec", {"text": "Build a widget."}, kind="spec"),
        _artifact("migration_plan", {"plan": "none needed"}),
    ])
    expected_titles = [s.title for s in render_sections(package)]

    text_output = render_text(package)
    for title in expected_titles:
        assert f"## {title}" in text_output

    out_path = tmp_path / "package.docx"
    render_docx(package, out_path)
    doc = Document(str(out_path))
    docx_titles = [p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    assert docx_titles == expected_titles
